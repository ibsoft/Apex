"""EDITOR tools: generate formatted Microsoft Word (.docx) and Excel (.xlsx)
documents and return signed, expiring download links."""
from __future__ import annotations

import base64
import io
import json
import mimetypes
import os
import re
import time
import traceback
import uuid
from pathlib import Path

import requests
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from flask import jsonify, send_file
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
from openpyxl import Workbook
from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference, ScatterChart
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from public_urls import public_url
from tools.base import Tool, ToolContext

DEFAULT_TABLE_STYLE = "Light Grid Accent 1"
EDITOR_SALT = "apex-editor-download-v1"


def _generated_dir(config, user_id: str) -> Path:
    base = Path(getattr(config, "DATA_DIR", Path(__file__).resolve().parents[1] / "data"))
    path = (base / "generated" / "editor" / str(user_id)).resolve()
    path.mkdir(parents=True, exist_ok=True)
    return path


def _safe_filename(name: str, default: str = "document", ext: str | None = None) -> str:
    name = str(name or default).strip()
    # Drop any path traversal attempts; keep only the basename.
    name = Path(name).name
    if not name:
        name = default
    if ext and not name.lower().endswith(ext.lower()):
        name = name + ext
    # Replace characters illegal in common filesystems.
    name = re.sub(r'[<>\:"/\\|?*]', "_", name)
    if not name or name.lower() == ext.lower():
        name = default + (ext or "")
    return name


def _cleanup_old_files(directory: Path, ttl_seconds: int):
    if not directory.exists():
        return
    cutoff = time.time() - ttl_seconds
    for item in directory.iterdir():
        try:
            if item.is_file() and item.stat().st_mtime < cutoff:
                item.unlink()
        except OSError:
            pass


def _fingerprint(info: os.stat_result) -> list:
    return [info.st_dev, info.st_ino, info.st_size, info.st_mtime_ns, info.st_ctime_ns]


def _signer(config):
    key = getattr(config, "SECRET_KEY", "dev-change-me")
    return URLSafeTimedSerializer(key, salt=EDITOR_SALT)


def _load_spec(args: dict, ctx: ToolContext) -> tuple[dict | None, str]:
    if not ctx.user_id:
        return None, "Sign in before creating documents."
    raw = args.get("document") or args.get("spec") or ""
    if not raw:
        return None, "Please provide a document specification as JSON."
    try:
        spec = json.loads(raw) if isinstance(raw, str) else dict(raw)
    except Exception as exc:
        return None, f"Invalid JSON document specification: {exc}"
    if not isinstance(spec, dict):
        return None, "Document specification must be a JSON object."
    return spec, ""


def _rgb_color(value: str) -> RGBColor:
    value = str(value).strip().lstrip("#")
    if len(value) == 3:
        value = "".join(c * 2 for c in value)
    if len(value) == 6:
        try:
            return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))
        except ValueError:
            pass
    return RGBColor(0, 0, 0)


def _word_align(align: str):
    mapping = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    return mapping.get(str(align).lower(), WD_ALIGN_PARAGRAPH.LEFT)


def _apply_word_run_format(run, fmt: dict):
    run.bold = bool(fmt.get("bold"))
    run.italic = bool(fmt.get("italic"))
    run.underline = bool(fmt.get("underline"))
    if fmt.get("font"):
        run.font.name = str(fmt["font"])
    if fmt.get("font_size"):
        try:
            run.font.size = Pt(float(fmt["font_size"]))
        except (TypeError, ValueError):
            pass
    color = fmt.get("color") or fmt.get("font_color")
    if color:
        run.font.color.rgb = _rgb_color(color)


def _fetch_image_data(source: str) -> bytes | None:
    source = str(source).strip()
    if not source:
        return None
    if source.startswith("data:"):
        m = re.match(r"data:image/[^;]+;base64,(.*)", source)
        if m:
            try:
                return base64.b64decode(m.group(1))
            except Exception:
                return None
        return None
    if source.lower().startswith(("http://", "https://")):
        try:
            resp = requests.get(source, timeout=30)
            resp.raise_for_status()
            return resp.content
        except Exception:
            return None
    return None


def _add_word_table(doc: Document, element: dict):
    headers = element.get("headers") or []
    rows = element.get("rows") or []
    if not rows and not headers:
        return
    cols = max(len(headers), max((len(r) for r in rows), default=0), 1)
    table = doc.add_table(rows=1, cols=cols)
    style = element.get("style") or DEFAULT_TABLE_STYLE
    try:
        table.style = style
    except Exception:
        try:
            table.style = DEFAULT_TABLE_STYLE
        except Exception:
            pass
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        if i < len(hdr_cells):
            hdr_cells[i].text = str(h)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            if i < len(cells):
                cells[i].text = "" if val is None else str(val)


def _add_word_image(doc: Document, element: dict):
    source = element.get("url") or element.get("source") or element.get("base64") or ""
    data = _fetch_image_data(source)
    if data is None:
        doc.add_paragraph(f"[Image not available: {str(source)[:80]}]")
        return
    try:
        pic = doc.add_picture(io.BytesIO(data))
        if element.get("width"):
            pic.width = Inches(float(element["width"]))
        if element.get("height"):
            pic.height = Inches(float(element["height"]))
    except Exception as exc:
        doc.add_paragraph(f"[Could not add image: {exc}]")
        return
    if element.get("caption"):
        p = doc.add_paragraph(str(element["caption"]))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p.runs:
            p.runs[0].italic = True


def _add_word_chart(doc: Document, element: dict):
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except Exception as exc:
        doc.add_paragraph(f"[Charts require matplotlib: {exc}]")
        return

    chart_type = str(element.get("chart_type", "bar")).lower()
    labels = element.get("labels") or []
    datasets = element.get("datasets") or []
    width = float(element.get("width", 6))
    height = float(element.get("height", 4))
    title = element.get("title") or ""

    fig, ax = plt.subplots(figsize=(width, height))
    if title:
        ax.set_title(title)

    for ds in datasets:
        values = ds.get("data") or []
        label = ds.get("label", "")
        color = ds.get("color")
        if chart_type == "line":
            ax.plot(labels[: len(values)], values, label=label or None, color=color)
        elif chart_type == "scatter":
            ax.scatter(range(len(values)), values, label=label or None, color=color)
        else:
            ax.bar(labels[: len(values)], values, label=label or None, color=color)

    if labels and chart_type not in ("scatter",):
        ax.set_xticks(range(len(labels)))
        ax.set_xticklabels(labels)

    if len(datasets) > 1:
        ax.legend()

    fig.tight_layout()
    buf = io.BytesIO()
    plt.savefig(buf, format="png", dpi=150)
    plt.close(fig)
    buf.seek(0)
    doc.add_picture(buf, width=Inches(width))


def _add_word_list(doc: Document, element: dict):
    items = element.get("items") or []
    ordered = bool(element.get("ordered", False))
    for i, item in enumerate(items, start=1):
        if ordered:
            doc.add_paragraph(f"{i}. {item}", style="List Number")
        else:
            doc.add_paragraph(str(item), style="List Bullet")


def _add_word_element(doc: Document, element: dict):
    if not isinstance(element, dict):
        doc.add_paragraph(str(element))
        return
    etype = str(element.get("type", "paragraph")).lower()
    if etype == "heading":
        level = min(max(int(element.get("level", 1)), 1), 9)
        doc.add_heading(str(element.get("text", "")), level=level)
    elif etype == "paragraph":
        style = element.get("style") or "Normal"
        p = doc.add_paragraph(str(element.get("text", "")), style=style)
        if element.get("align"):
            p.alignment = _word_align(element["align"])
        if p.runs:
            _apply_word_run_format(p.runs[0], element)
    elif etype == "table":
        _add_word_table(doc, element)
    elif etype == "image":
        _add_word_image(doc, element)
    elif etype == "chart":
        _add_word_chart(doc, element)
    elif etype == "page_break":
        doc.add_page_break()
    elif etype == "list":
        _add_word_list(doc, element)
    else:
        doc.add_paragraph(str(element.get("text", "")))


def _build_word(spec: dict, path: Path):
    doc = Document()
    # Sensible default body font.
    try:
        style = doc.styles["Normal"]
        if style.font:
            style.font.name = "Calibri"
            style.font.size = Pt(11)
    except Exception:
        pass

    title = spec.get("title")
    if title:
        try:
            p = doc.add_heading(str(title), level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:
            doc.add_paragraph(str(title))

    sections = spec.get("sections") or []
    if not isinstance(sections, list):
        raise ValueError("'sections' must be a list.")
    for element in sections:
        _add_word_element(doc, element)

    doc.save(str(path))


def _apply_excel_cell_style(cell, style: dict):
    if not isinstance(style, dict):
        return
    font_kwargs = {}
    if style.get("bold"):
        font_kwargs["bold"] = True
    if style.get("italic"):
        font_kwargs["italic"] = True
    if style.get("underline"):
        font_kwargs["underline"] = "single"
    if style.get("font_size"):
        try:
            font_kwargs["size"] = float(style["font_size"])
        except (TypeError, ValueError):
            pass
    if style.get("font_color"):
        font_kwargs["color"] = str(style["font_color"]).lstrip("#").upper()
    if font_kwargs:
        cell.font = Font(**font_kwargs)

    if style.get("fill_color"):
        fill = str(style["fill_color"]).lstrip("#").upper()
        cell.fill = PatternFill(start_color=fill, end_color=fill, fill_type="solid")

    if style.get("number_format"):
        cell.number_format = str(style["number_format"])

    align = style.get("align")
    if align:
        cell.alignment = Alignment(horizontal=str(align).lower())

    border = style.get("border")
    if border:
        side = Side(style=str(border).lower())
        cell.border = Border(left=side, right=side, top=side, bottom=side)


def _excel_col_to_idx(col: str) -> int:
    idx = 0
    for c in col.upper():
        idx = idx * 26 + (ord(c) - ord("A") + 1)
    return idx


def _add_excel_chart(ws, chart_spec: dict, data: list[list]):
    chart_type = str(chart_spec.get("type", "bar")).lower()
    title = chart_spec.get("title", "")
    data_range = chart_spec.get("data_range")
    if not data_range and data:
        last_row = len(data)
        last_col = max((len(row) for row in data), default=1)
        data_range = f"A1:{get_column_letter(last_col)}{last_row}"

    if not data_range:
        return
    m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", data_range.upper())
    if not m:
        return
    c1, r1, c2, r2 = m.groups()
    min_col = _excel_col_to_idx(c1)
    min_row = int(r1)
    max_col = _excel_col_to_idx(c2)
    max_row = int(r2)

    if chart_type == "bar":
        chart = BarChart()
    elif chart_type == "line":
        chart = LineChart()
    elif chart_type == "pie":
        chart = PieChart()
    elif chart_type == "area":
        chart = AreaChart()
    elif chart_type == "scatter":
        chart = ScatterChart()
    else:
        chart = BarChart()

    chart.title = title
    categories_first = bool(chart_spec.get("categories_from_first_column", True))
    series_first = bool(chart_spec.get("series_from_first_row", True))

    data_min_col = min_col + (1 if categories_first else 0)
    data_min_row = min_row + (1 if series_first else 0)
    if data_min_col <= max_col and data_min_row <= max_row:
        data_ref = Reference(ws, min_col=data_min_col, min_row=data_min_row,
                             max_col=max_col, max_row=max_row)
        chart.add_data(data_ref, titles_from_data=series_first)

    if categories_first:
        cats = Reference(ws, min_col=min_col, min_row=min_row + 1, max_row=max_row)
        chart.set_categories(cats)

    if chart_spec.get("x_axis_title"):
        chart.x_axis.title = chart_spec["x_axis_title"]
    if chart_spec.get("y_axis_title"):
        chart.y_axis.title = chart_spec["y_axis_title"]

    pos = chart_spec.get("position") or "E5"
    ws.add_chart(chart, pos)


def _fill_excel_sheet(ws, spec: dict):
    data = spec.get("data") or []
    for r_idx, row in enumerate(data, start=1):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            if r_idx == 1:
                header_style = spec.get("header_style") or {"bold": True, "fill_color": "DDDDDD"}
                _apply_excel_cell_style(cell, header_style)
            else:
                cell_style = spec.get("cell_style")
                if cell_style:
                    _apply_excel_cell_style(cell, cell_style)

    formulas = spec.get("formulas") or {}
    for cell_ref, formula in formulas.items():
        ws[str(cell_ref)] = formula

    widths = spec.get("column_widths") or {}
    for col, width in widths.items():
        col_letter = get_column_letter(int(col)) if isinstance(col, int) else str(col).upper()
        ws.column_dimensions[col_letter].width = float(width)

    if spec.get("freeze_header", True) and len(data) > 0:
        ws.freeze_panes = "A2"

    for chart_spec in spec.get("charts") or []:
        _add_excel_chart(ws, chart_spec, data)


def _build_excel(spec: dict, path: Path):
    sheets = spec.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        raise ValueError("'sheets' must be a non-empty list.")
    wb = Workbook()
    for i, sheet_spec in enumerate(sheets):
        title = str(sheet_spec.get("name", "Sheet1" if i == 0 else f"Sheet{i + 1}"))
        if i == 0:
            ws = wb.active
            ws.title = title[:31]
        else:
            ws = wb.create_sheet(title=title[:31])
        _fill_excel_sheet(ws, sheet_spec)
    wb.save(str(path))


def _create_document(args: dict, ctx: ToolContext, config, ext: str, builder):
    spec, err = _load_spec(args, ctx)
    if err:
        return err
    if not getattr(config, "EDITOR_ENABLED", True):
        return "Document editor is disabled in configuration."

    filename = _safe_filename(spec.get("filename"), default="document", ext=ext)
    directory = _generated_dir(config, ctx.user_id)
    ttl = getattr(config, "EDITOR_FILE_TTL_SECONDS", 3600)
    _cleanup_old_files(directory, ttl)

    path = directory / f"{uuid.uuid4().hex}_{filename}"
    try:
        builder(spec, path)
    except Exception as exc:
        return f"Could not generate document: {exc}\n{traceback.format_exc()[-500:]}"

    try:
        info = path.stat()
        token = _signer(config).dumps(
            {"user": str(ctx.user_id), "path": str(path),
             "fingerprint": _fingerprint(info), "filename": filename}
        )
        url = public_url(config, f"/api/editor/download/{token}")
    except Exception as exc:
        return f"Document generated but could not create download link: {exc}"

    return json.dumps(
        {"filename": filename, "download_url": url, "expires_in_seconds": ttl},
        ensure_ascii=False,
    )


def build_editor_tools(config) -> list[Tool]:
    def t_editor_create_word(args: dict, ctx: ToolContext):
        return _create_document(args, ctx, config, ".docx", _build_word)

    def t_editor_create_excel(args: dict, ctx: ToolContext):
        return _create_document(args, ctx, config, ".xlsx", _build_excel)

    word_schema = {
        "type": "object",
        "properties": {
            "document": {
                "type": "string",
                "description": "JSON document specification with keys: filename, title, sections. "
                               "Sections support type: heading, paragraph, table, image, chart, list, page_break.",
            },
        },
        "required": ["document"],
        "additionalProperties": False,
    }

    excel_schema = {
        "type": "object",
        "properties": {
            "document": {
                "type": "string",
                "description": "JSON workbook specification with keys: filename, sheets. "
                               "Each sheet has name, data, column_widths, header_style, cell_style, formulas, charts.",
            },
        },
        "required": ["document"],
        "additionalProperties": False,
    }

    return [
        Tool(
            "editor_create_word",
            "Create a formatted Microsoft Word document (.docx) and return a signed download link.",
            word_schema,
            t_editor_create_word,
        ),
        Tool(
            "editor_create_excel",
            "Create a formatted Microsoft Excel workbook (.xlsx) with tables, formatting and charts, and return a signed download link.",
            excel_schema,
            t_editor_create_excel,
        ),
    ]


def register_editor_routes(app, require_user, config):
    @app.get("/api/editor/download/<token>")
    def download_editor_file(token):
        user = require_user()
        if user is None:
            return jsonify({"error": "Sign in to download this file."}), 401
        ttl = getattr(config, "EDITOR_FILE_TTL_SECONDS", 3600)
        try:
            ticket = _signer(config).loads(token, max_age=ttl)
        except SignatureExpired:
            return jsonify({"error": "Download link expired. Generate the document again."}), 410
        except BadSignature:
            return jsonify({"error": "Invalid download link."}), 404
        if not isinstance(ticket, dict) or ticket.get("user") != str(user["id"]):
            return jsonify({"error": "This download link belongs to a different user."}), 403
        try:
            path = Path(ticket["path"]).resolve()
            base_dir = (
                Path(getattr(config, "DATA_DIR", Path(__file__).resolve().parents[1] / "data"))
                / "generated"
                / "editor"
            ).resolve()
            if not path.is_relative_to(base_dir):
                return jsonify({"error": "Invalid file location."}), 403
            if not path.is_file():
                return jsonify({"error": "File is unavailable."}), 404
            info = path.stat()
        except (OSError, ValueError, KeyError, TypeError):
            return jsonify({"error": "File is unavailable."}), 404
        if _fingerprint(info) != ticket.get("fingerprint"):
            return jsonify({"error": "File changed. Generate it again."}), 409
        try:
            download_name = ticket.get("filename") or path.name
            mimetype = mimetypes.guess_type(download_name)[0] or "application/octet-stream"
            response = send_file(
                path, as_attachment=True, download_name=download_name,
                mimetype=mimetype, conditional=False, max_age=0,
            )
        except Exception:
            return jsonify({"error": "Could not send file."}), 500
        response.headers["Cache-Control"] = "private, no-store"
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["Referrer-Policy"] = "no-referrer"
        return response
