"""Document ingestion for memory: extract text and chunk it for embedding."""
from __future__ import annotations

import io
import os
import re
import time
import uuid
from pathlib import Path

_DOCX_MIMES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.ms-word.document.macroEnabled.12",
}
_XLSX_MIMES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel.sheet.macroEnabled.12",
}


def _guess_mimetype(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    return {
        ".txt": "text/plain",
        ".md": "text/markdown",
        ".markdown": "text/markdown",
        ".py": "text/x-python",
        ".js": "text/javascript",
        ".ts": "text/typescript",
        ".jsx": "text/jsx",
        ".tsx": "text/tsx",
        ".json": "application/json",
        ".csv": "text/csv",
        ".html": "text/html",
        ".htm": "text/html",
        ".xml": "text/xml",
        ".yaml": "text/yaml",
        ".yml": "text/yaml",
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".docm": "application/vnd.ms-word.document.macroEnabled.12",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
    }.get(ext, "application/octet-stream")


def _docx_text(data: bytes) -> str:
    """Extract paragraphs and tables from a .docx in document order (UTF-8)."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(io.BytesIO(data))
    parts: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "").lower()
            if style.startswith("list"):
                parts.append(f"• {text}")
            else:
                parts.append(text)
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
    return "\n".join(parts)


def _xlsx_text(data: bytes) -> str:
    """Extract every sheet's cell values from a .xlsx as read, text rows (UTF-8)."""
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        parts: list[str] = []
        for sheet in workbook.worksheets:
            parts.append(f"# {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = [("" if value is None else str(value)).strip() for value in row]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    finally:
        workbook.close()


def extract_text(filename: str, data: bytes) -> str:
    """Best-effort text extraction from common document formats."""
    mimetype = _guess_mimetype(filename)

    if mimetype in _DOCX_MIMES:
        try:
            return _docx_text(data)
        except ImportError as exc:
            raise RuntimeError(
                "Word documents require python-docx (pip install python-docx)"
            ) from exc

    if mimetype in _XLSX_MIMES:
        try:
            return _xlsx_text(data)
        except ImportError as exc:
            raise RuntimeError(
                "Excel documents require openpyxl (pip install openpyxl)"
            ) from exc

    if mimetype == "application/pdf":
        try:
            from PyPDF2 import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF support requires PyPDF2 (pip install PyPDF2)") from exc
        reader = PdfReader(io.BytesIO(data))
        parts = []
        for page in reader.pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                pass
        return "\n\n".join(parts)

    # Everything else is treated as UTF-8 text.
    # A few formats are common enough that ignoring non-text bytes is fine.
    text = data.decode("utf-8", errors="ignore")
    if mimetype == "text/html":
        text = re.sub(r"<[^>]+>", " ", text)
        text = re.sub(r"\s+", " ", text)
    return text


# ---------------------------------------------------------------------------
# Uploaded-file persistence. Memory ingests the extracted *text*, but sending a
# document on (e.g. by the EMAIL skill) needs the original bytes on disk. The
# upload route calls store_upload so the model can reference the returned path
# with --attach instead of having to re-create the file.

_UPLOAD_TTL_DAYS = 7   # far longer than any attach/email turn needs
_UPLOAD_ROOT_NAME = "uploads"


def _upload_root(config, user_id: str) -> Path:
    data_dir = Path(getattr(config, "DATA_DIR", Path(__file__).resolve().parents[1] / "data"))
    root = data_dir / _UPLOAD_ROOT_NAME / str(user_id)
    root.mkdir(parents=True, exist_ok=True)
    return root


def store_upload(config, user_id: str, filename: str, data: bytes) -> Path:
    """Persist one uploaded file for later attachment, owner-scoped (0600).

    The stored name is ``<uuid>_<basename>`` so a re-upload of an identical
    filename cannot collide and no path may escape the user's directory.
    """
    if not data:
        raise ValueError("empty file")
    safe = re.sub(r"[^\w.\-]", "_", Path(filename or "file").name) or "file"
    safe = safe.strip(".")[:180] or "file"
    target = _upload_root(config, user_id) / f"{uuid.uuid4().hex[:12]}_{safe}"
    tmp = target.with_name(target.name + ".tmp")
    tmp.write_bytes(data)
    os.chmod(tmp, 0o600)
    os.replace(tmp, target)
    return target


def sweep_uploads(config, user_id: str) -> None:
    """Drop uploaded files older than the TTL so they cannot accumulate."""
    root = _upload_root(config, user_id)
    cutoff = time.time() - _UPLOAD_TTL_DAYS * 86400
    for entry in root.iterdir():
        try:
            if entry.is_file() and not entry.name.endswith(".tmp") \
                    and entry.stat().st_mtime < cutoff:
                entry.unlink()
        except OSError:
            continue


def chunk_text(
    text: str,
    chunk_size: int = 800,
    overlap: int = 100,
) -> list[str]:
    """Split text into overlapping chunks.

    Splits on paragraphs first, then sentences, then words, so chunks stay
    semantically coherent while staying under ``chunk_size`` characters.
    """
    text = re.sub(r"\r\n?", "\n", text).strip()
    if not text:
        return []
    if len(text) <= chunk_size:
        return [text]

    # 1. Split on paragraphs.
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]

    chunks: list[str] = []
    current = ""

    def _flush():
        nonlocal current
        if current:
            chunks.append(current.strip())
            current = ""

    def _split_by_sentences(block: str) -> list[str]:
        # Simple sentence split on punctuation followed by space or newline.
        sentences = re.split(r"(?<=[.!?])\s+", block)
        return [s.strip() for s in sentences if s.strip()]

    def _split_by_words(block: str) -> list[str]:
        words = block.split()
        out = []
        buf: list[str] = []
        length = 0
        for w in words:
            if length + len(w) + 1 > chunk_size and buf:
                out.append(" ".join(buf))
                # carry over overlap words
                carry = []
                carry_len = 0
                for cw in reversed(buf):
                    if carry_len + len(cw) + 1 > overlap:
                        break
                    carry.insert(0, cw)
                    carry_len += len(cw) + 1
                buf = carry
                length = carry_len
            buf.append(w)
            length += len(w) + 1
        if buf:
            out.append(" ".join(buf))
        return out

    for para in paragraphs:
        if len(para) > chunk_size:
            _flush()
            sentences = _split_by_sentences(para)
            sent_buf = ""
            for sent in sentences:
                if len(sent_buf) + len(sent) + 1 > chunk_size and sent_buf:
                    # sentence buffer too big; split it by words
                    chunks.extend(_split_by_words(sent_buf.strip()))
                    sent_buf = ""
                sent_buf += (" " if sent_buf else "") + sent
            if sent_buf:
                if len(sent_buf) > chunk_size:
                    chunks.extend(_split_by_words(sent_buf.strip()))
                else:
                    chunks.append(sent_buf.strip())
            continue

        if current and len(current) + len(para) + 2 > chunk_size:
            _flush()
        current += ("\n\n" if current else "") + para

    _flush()

    # Apply overlap between adjacent chunks.
    if overlap <= 0 or len(chunks) < 2:
        return chunks

    overlapped: list[str] = [chunks[0]]
    for i in range(1, len(chunks)):
        prev = chunks[i - 1]
        curr = chunks[i]
        prefix = prev[-overlap:].lstrip()
        combined = (prefix + " " + curr).strip()
        if len(combined) <= chunk_size:
            overlapped.append(combined)
        else:
            overlapped.append(curr)
    return overlapped
