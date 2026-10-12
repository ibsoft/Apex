"""Integration test for document upload into memory."""
from __future__ import annotations

import io
import os
import tempfile
import unittest
from pathlib import Path


class MemoryUploadTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.orig_data_dir = os.environ.get("DATA_DIR")
        os.environ["DEV_MODE"] = "true"
        os.environ["DATA_DIR"] = self.tmp.name

        from app import create_app
        from config import config
        from memory import store

        self.config = config
        config.DATA_DIR = Path(self.tmp.name)
        config.CHROMA_DIR = str(Path(self.tmp.name) / "chroma")
        config.MEMORY_ENABLED = True
        store.register_memory_store(None)  # force a fresh store on next call

        # This test is about uploads, not sign-in, so the test client is given
        # an identity outright. DEV_MODE alone no longer does it: dev auto-login
        # is ignored while system login is on, which used to leave every request
        # here a 401 whenever .env happened to be clean.
        self.orig_auth_flags = (
            config.SYSTEM_LOGIN_ENABLED,
            config.DEV_AUTO_LOGIN,
            config.DEV_AUTO_LOGIN_WITH_SYSTEM_LOGIN,
        )
        config.SYSTEM_LOGIN_ENABLED = False
        config.DEV_AUTO_LOGIN = "uploader"
        config.DEV_AUTO_LOGIN_WITH_SYSTEM_LOGIN = False

        self.app = create_app()
        self.client = self.app.test_client()

    def tearDown(self):
        if self.orig_data_dir is None:
            os.environ.pop("DATA_DIR", None)
        else:
            os.environ["DATA_DIR"] = self.orig_data_dir
        (system_login_enabled, dev_auto_login, dev_with_system_login) = self.orig_auth_flags
        self.config.SYSTEM_LOGIN_ENABLED = system_login_enabled
        self.config.DEV_AUTO_LOGIN = dev_auto_login
        self.config.DEV_AUTO_LOGIN_WITH_SYSTEM_LOGIN = dev_with_system_login
        self.tmp.cleanup()

    def test_upload_text_file(self):
        data = {"files": (io.BytesIO(b"hello world\n\nthis is a test document"), "test.txt")}
        resp = self.client.post(
            "/api/memory/upload",
            data=data,
            content_type="multipart/form-data",
        )
        self.assertEqual(resp.status_code, 200)
        body = resp.get_json()
        self.assertTrue(body["ok"])
        self.assertGreater(body["total"], 0)
        self.assertEqual(body["files"][0]["filename"], "test.txt")
        # A bounded preview travels back so the attaching turn can answer from
        # the document without a recall round-trip.
        self.assertIn("hello world", body["files"][0]["preview"])
        # The original bytes are persisted so a skill (EMAIL --attach) can send
        # the real document; the path points inside the uploader's own dir.
        saved = Path(body["files"][0]["path"])
        self.assertEqual(saved.read_bytes(), b"hello world\n\nthis is a test document")
        self.assertTrue(saved.is_relative_to(Path(self.tmp.name) / "uploads" / "uploader"))

    def test_upload_docx(self):
        from docx import Document

        document = Document()
        document.add_paragraph("A Word document uploaded to memory.")
        buf = io.BytesIO()
        document.save(buf)
        data = {"files": (io.BytesIO(buf.getvalue()), "memo.docx")}
        resp = self.client.post(
            "/api/memory/upload",
            data=data,
            content_type="multipart/form-data",
        )
        self.assertEqual(resp.status_code, 200)
        body = resp.get_json()
        self.assertGreater(body["total"], 0)
        self.assertIn("A Word document uploaded to memory.", body["files"][0]["preview"])
        self.assertTrue(body["files"][0]["path"])

    def test_upload_xlsx(self):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(["Name", "Count"])
        sheet.append(["alpha", 1])
        buf = io.BytesIO()
        workbook.save(buf)
        data = {"files": (io.BytesIO(buf.getvalue()), "table.xlsx")}
        resp = self.client.post(
            "/api/memory/upload",
            data=data,
            content_type="multipart/form-data",
        )
        self.assertEqual(resp.status_code, 200)
        body = resp.get_json()
        self.assertGreater(body["total"], 0)
        preview = body["files"][0]["preview"]
        self.assertIn("Name", preview)
        self.assertIn("alpha", preview)
        self.assertTrue(body["files"][0]["path"])

    def test_upload_preview_is_bounded(self):
        text = b"x" * 5000
        data = {"files": (io.BytesIO(text), "long.txt")}
        resp = self.client.post(
            "/api/memory/upload",
            data=data,
            content_type="multipart/form-data",
        )
        self.assertEqual(resp.status_code, 200)
        preview = resp.get_json()["files"][0]["preview"]
        self.assertEqual(len(preview), 2000)

    def test_upload_rejects_no_files(self):
        resp = self.client.post("/api/memory/upload", data={})
        self.assertEqual(resp.status_code, 400)


if __name__ == "__main__":
    unittest.main()
