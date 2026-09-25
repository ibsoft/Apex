"""Tests for the window-manager preview renderer (docx/xlsx/text -> HTML)."""
from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from flask import Flask

from tools.base import ToolContext
from tools.editor_tools import build_editor_tools, register_editor_routes
from tools.preview_tools import register_preview_routes

# A real 8x8 PNG so python-docx (which validates PNG chunk structure) accepts it.
def _png_bytes() -> bytes:
    import io

    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (8, 8), (200, 40, 40)).save(buffer, format="PNG")
    return buffer.getvalue()


PNG_BYTES = _png_bytes()


class PreviewRenderTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.data_dir = Path(self.temporary.name)
        self.images_dir = self.data_dir / "images"
        self.images_dir.mkdir()
        self.vault = self.data_dir / "vault"
        self.vault.mkdir()
        self.config = SimpleNamespace(
            DATA_DIR=self.data_dir,
            SECRET_KEY="test-secret",
            BASE_URL="http://localhost",
            EDITOR_ENABLED=True,
            EDITOR_FILE_TTL_SECONDS=3600,
            IMAGES_DIR=str(self.images_dir),
            OBSIDIAN_VAULT_PATH=str(self.vault),
        )
        self.app = Flask(__name__)
        self.app.secret_key = self.config.SECRET_KEY

        def require_user():
            from flask import session

            return {"id": session["user_id"]} if session.get("user_id") else None

        register_editor_routes(self.app, require_user, self.config)
        register_preview_routes(self.app, require_user, self.config)
        self.client = self.app.test_client()
        self.login("alice")
        self.word_tool, self.excel_tool = build_editor_tools(self.config)
        self.ctx = ToolContext(user_id="alice")

    def tearDown(self):
        self.temporary.cleanup()

    def login(self, user):
        with self.client.session_transaction() as sess:
            sess.clear()
            if user:
                sess["user_id"] = user

    def render(self, url, user="alice"):
        self.login(user)
        return self.client.get("/api/preview/render", query_string={"url": url})

    def make_editor_url(self, spec, tool=None):
        tool = tool or self.word_tool
        return json.loads(tool.call({"document": json.dumps(spec)}, self.ctx))["download_url"]

    def sign_images_link(self, path, user=None):
        from tools.image_browser import _fingerprint, _signer

        token = _signer(self.config).dumps(
            {"user": user, "path": str(path), "fingerprint": _fingerprint(path.stat())}
        )
        return f"/api/images/file/{token}"

    # ---- EDITOR documents ---------------------------------------------------

    def test_word_document_renders_as_html(self):
        url = self.make_editor_url({
            "filename": "report.docx",
            "sections": [
                {"type": "heading", "level": 1, "text": "Overview"},
                {"type": "paragraph", "text": "Quarterly numbers enclosed."},
                {"type": "table", "headers": ["Item", "Value"], "rows": [["A", 1], ["B", 2]]},
            ],
        })
        response = self.client.get("/api/preview/render", query_string={"url": url})
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/html", response.headers["Content-Type"])
        text = response.get_data(as_text=True)
        self.assertIn("Quarterly numbers enclosed.", text)
        self.assertIn("<h1>Overview</h1>", text)
        self.assertIn("<th>Item</th>", text)

    def test_excel_workbook_renders_as_html(self):
        url = self.make_editor_url({
            "filename": "data.xlsx",
            "sheets": [{"name": "Sales", "data": [["Product", "Q1"], ["Widget", 120]]}],
        }, tool=self.excel_tool)
        response = self.client.get("/api/preview/render", query_string={"url": url})
        self.assertEqual(response.status_code, 200)
        text = response.get_data(as_text=True)
        self.assertIn("text/html", response.headers["Content-Type"])
        self.assertIn("Sales", text)
        self.assertIn("Widget", text)
        self.assertIn("<th>Product</th>", text)

    def test_docx_preview_preserves_lists(self):
        url = self.make_editor_url({
            "filename": "notes.docx",
            "sections": [
                {"type": "list", "items": ["alpha", "beta"]},
                {"type": "paragraph", "text": "end"},
            ],
        })
        text = self.client.get("/api/preview/render", query_string={"url": url}).get_data(as_text=True)
        self.assertIn("<li>alpha</li>", text)
        self.assertIn("<li>beta</li>", text)
        self.assertIn("<p>end</p>", text)

    def test_docx_preview_embeds_generated_images(self):
        import base64

        url = self.make_editor_url({
            "filename": "withimg.docx",
            "sections": [
                {"type": "image", "base64": "data:image/png;base64," + base64.b64encode(PNG_BYTES).decode("ascii")},
            ],
        })
        text = self.client.get("/api/preview/render", query_string={"url": url}).get_data(as_text=True)
        self.assertIn("data:image", text)

    # ---- PPTX ---------------------------------------------------------------

    def make_pptx_bytes(self):
        import io

        from pptx import Presentation

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "The Moon"
        slide.placeholders[1].text_frame.text = "Earth's natural satellite"
        table = slide.shapes.add_table(2, 2, 0, 0, 200, 80).table
        table.cell(0, 0).text = "Radius"
        table.cell(0, 1).text = "1,737 km"
        table.cell(1, 0).text = "Orbit"
        table.cell(1, 1).text = "27.3 days"
        buffer = io.BytesIO()
        prs.save(buffer)
        return buffer.getvalue()

    def test_pptx_renders_as_html(self):
        deck = self.images_dir / "moon_deck.pptx"
        deck.write_bytes(self.make_pptx_bytes())
        response = self.client.get(
            "/api/preview/render",
            query_string={"url": self.sign_images_link(deck)},
        )
        self.assertEqual(response.status_code, 200)
        text = response.get_data(as_text=True)
        self.assertIn("text/html", response.headers["Content-Type"])
        self.assertIn("Slide 1", text)
        self.assertIn("The Moon", text)
        self.assertIn("1,737 km", text)

    def test_pptx_kind_endpoint(self):
        deck = self.images_dir / "moon_deck.pptx"
        deck.write_bytes(self.make_pptx_bytes())
        kind = self.client.get(
            "/api/preview/kind",
            query_string={"url": self.sign_images_link(deck)},
        ).get_json()
        self.assertEqual(kind["kind"], "pptx")

    # ---- auth / link validation --------------------------------------------

    def test_requires_sign_in(self):
        url = self.make_editor_url({
            "filename": "x.docx",
            "sections": [{"type": "paragraph", "text": "Hi"}],
        })
        self.login(None)
        self.assertEqual(self.client.get("/api/preview/render", query_string={"url": url}).status_code, 401)

    def test_owner_is_enforced(self):
        url = self.make_editor_url({
            "filename": "x.docx",
            "sections": [{"type": "paragraph", "text": "Hi"}],
        })
        self.assertEqual(self.render(url, user="bob").status_code, 403)

    def test_tampered_token_rejected(self):
        url = self.make_editor_url({
            "filename": "x.docx",
            "sections": [{"type": "paragraph", "text": "Hi"}],
        })
        self.assertEqual(
            self.client.get("/api/preview/render", query_string={"url": url + "x"}).status_code,
            404,
        )

    def test_unknown_url_shape_rejected(self):
        response = self.client.get(
            "/api/preview/render",
            query_string={"url": "https://example.com/remote.png"},
        )
        self.assertEqual(response.status_code, 400)

    def test_missing_url_rejected(self):
        self.assertEqual(self.client.get("/api/preview/render").status_code, 400)

    # ---- image / text / unsupported ----------------------------------------

    def test_image_signed_link_serves_inline(self):
        img = self.images_dir / "snap.png"
        img.write_bytes(PNG_BYTES)
        response = self.client.get(
            "/api/preview/render",
            query_string={"url": f"http://localhost{self.sign_images_link(img)}"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data, PNG_BYTES)
        self.assertIn("image/png", response.headers["Content-Type"])

    def test_image_link_belongs_to_other_user_rejected(self):
        img = self.images_dir / "snap.png"
        img.write_bytes(PNG_BYTES)
        response = self.client.get(
            "/api/preview/render",
            query_string={"url": self.sign_images_link(img, user="bob")},
        )
        self.assertEqual(response.status_code, 403)

    def test_text_file_renders_as_html(self):
        note = self.images_dir / "readme.txt"
        note.write_bytes("Hello window!\nSecond line".encode("utf-8"))
        response = self.client.get(
            "/api/preview/render",
            query_string={"url": self.sign_images_link(note)},
        )
        self.assertEqual(response.status_code, 200)
        text = response.get_data(as_text=True)
        self.assertIn("<pre>", text)
        self.assertIn("Hello window!", text)

    def test_unsupported_type_returns_415(self):
        blob = self.images_dir / "archive.zip"
        blob.write_bytes(b"PK")
        response = self.client.get(
            "/api/preview/render",
            query_string={"url": self.sign_images_link(blob)},
        )
        self.assertEqual(response.status_code, 415)

    # ---- Obsidian vault -----------------------------------------------------

    def test_obsidian_attachment_inline(self):
        img = self.vault / "attach" / "photo.png"
        img.parent.mkdir(parents=True)
        img.write_bytes(PNG_BYTES)
        response = self.client.get(
            "/api/preview/render",
            query_string={"url": "/api/obsidian/file?path=attach/photo.png"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data, PNG_BYTES)

    def test_obsidian_note_file_is_not_an_attachment(self):
        note = self.vault / "note.md"
        note.write_text("# Hi")
        response = self.client.get(
            "/api/preview/render",
            query_string={"url": "/api/obsidian/file?path=note.md"},
        )
        self.assertEqual(response.status_code, 404)

    def test_obsidian_path_traversal_rejected(self):
        response = self.client.get(
            "/api/preview/render",
            query_string={"url": "/api/obsidian/file?path=../outside.png"},
        )
        self.assertEqual(response.status_code, 400)


if __name__ == "__main__":
    unittest.main()