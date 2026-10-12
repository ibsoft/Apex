"""Tests for document ingestion helpers."""
from __future__ import annotations

import io
import os
import unittest
from pathlib import Path

from memory.documents import chunk_text, extract_text, store_upload, sweep_uploads


def _make_docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Quarterly Report", level=1)
    document.add_paragraph("Revenue grew by 12% year over year.")
    document.add_heading("Table", level=2)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Widgets"
    table.cell(1, 1).text = "42"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _make_xlsx() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["Item", "Qty"])
    sheet.append(["Widget", 3])
    sheet.append(["Gadget", 7])
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


class ExtractTextTests(unittest.TestCase):
    def test_extract_plain_text(self):
        text = extract_text("notes.txt", b"hello world\nsecond line")
        self.assertIn("hello world", text)

    def test_extract_json(self):
        text = extract_text("data.json", b'{"key": "value"}')
        self.assertIn('"key": "value"', text)

    def test_extract_docx_paragraphs_and_tables(self):
        text = extract_text("report.docx", _make_docx())
        self.assertIn("Quarterly Report", text)
        self.assertIn("Revenue grew by 12% year over year.", text)
        self.assertIn("Item", text)
        self.assertIn("Widgets", text)
        self.assertIn("42", text)

    def test_extract_xlsx_every_sheet(self):
        text = extract_text("sales.xlsx", _make_xlsx())
        self.assertIn("Sales", text)
        self.assertIn("Item", text)
        self.assertIn("Widget", text)
        self.assertIn("Gadget", text)


class StoreUploadTests(unittest.TestCase):
    def setUp(self):
        import tempfile
        from types import SimpleNamespace

        from pathlib import Path

        self.tmp = tempfile.TemporaryDirectory()
        self.config = SimpleNamespace(DATA_DIR=Path(self.tmp.name))

    def tearDown(self):
        self.tmp.cleanup()

    def test_store_upload_round_trips_owner_scoped_bytes(self):
        target = store_upload(self.config, "alice", "quote.docx", b"PDF-ish bytes")
        root = self.config.DATA_DIR / "uploads" / "alice"
        self.assertTrue(target.is_relative_to(root))
        self.assertTrue(target.is_file())
        self.assertEqual(target.read_bytes(), b"PDF-ish bytes")

    def test_store_upload_sanitizes_and_uniquifies_names(self):
        first = store_upload(self.config, "alice", "../../etc/passwd", b"a")
        second = store_upload(self.config, "alice", "../../etc/passwd", b"b")
        self.assertEqual(first.parent, second.parent)
        self.assertNotEqual(first.name, second.name)
        self.assertEqual(Path(first.name).name, first.name)  # no traversal

    def test_sweep_removes_only_old_files(self):
        old = store_upload(self.config, "alice", "old.txt", b"old")
        fresh = store_upload(self.config, "alice", "new.txt", b"new")
        os.utime(old, (1, 1))

        sweep_uploads(self.config, "alice")

        self.assertFalse(old.exists())
        self.assertTrue(fresh.exists())


class ChunkTextTests(unittest.TestCase):
    def test_no_chunking_for_short_text(self):
        chunks = chunk_text("short text", chunk_size=100)
        self.assertEqual(chunks, ["short text"])

    def test_chunking_large_text(self):
        text = "\n\n".join([f"Paragraph {i} with some words here." for i in range(50)])
        chunks = chunk_text(text, chunk_size=200, overlap=20)
        self.assertGreater(len(chunks), 1)
        for c in chunks:
            self.assertLessEqual(len(c), 200)

    def test_empty_text(self):
        self.assertEqual(chunk_text(""), [])


if __name__ == "__main__":
    unittest.main()